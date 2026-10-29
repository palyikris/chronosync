"""Prepare a tenant-scoped, repeatable price seed from an indexed-price workbook.

Read-only input processing; writes SQL, a manifest, and CSV review files. It never
connects to a database. The snapshot must contain clients and projects exported
for the requested tenant. Requires openpyxl for reading the workbook.
"""
import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from uuid import UUID

import openpyxl

HEADERS = ['Ügyfél neve', 'Ügyfélkód', 'Projekt neve', 'Díjazása',
           '2026. évi díjak', 'Egység', 'Deviza',
           'Díjemelés szerepel a szerződésben', 'Alapja az emelésnek',
           'Melléklet', 'PO szám']
QUANTUM = Decimal('0.0001')


def normalize(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(value or '')).strip()).casefold()


def money(value):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError(f'Expected a numeric fee: {value!r}')
    result = Decimal(str(value)).quantize(QUANTUM, rounding=ROUND_HALF_UP)
    if not result.is_finite() or result < 0 or result >= Decimal('100000000000000'):
        raise ValueError(f'Invalid fee: {value!r}')
    return str(result)


def indexed_fee(sheet, cached_sheet, row):
    formula = sheet.cell(row, 5).value
    cached = cached_sheet.cell(row, 5).value
    if not isinstance(formula, str) or not formula.startswith('='):
        return money(formula)
    direct = re.fullmatch(r'=D(\d+)\*(1\.0(?:24|4|44))', formula)
    inflation = re.fullmatch(r'=\+D(\d+)\*\(100%\+\$F\$([12])\)', formula)
    if direct:
        ref, factor = int(direct[1]), Decimal(direct[2])
    elif inflation:
        ref = int(inflation[1])
        factor = Decimal(1) + Decimal(str(sheet.cell(int(inflation[2]), 6).value))
    else:
        raise ValueError(f'Unsupported formula in E{row}: {formula}')
    if ref != row:
        raise ValueError(f'Unexpected cross-row price reference in E{row}')
    result = money(Decimal(str(sheet.cell(row, 4).value)) * factor)
    if cached is None or abs(Decimal(result) - Decimal(str(cached))) > QUANTUM:
        raise ValueError(f'Formula/cache mismatch in E{row}')
    return result


def prepare(source, snapshot, company_id, year):
    company_id = str(UUID(company_id))
    if any(r['company_id'] != company_id for key in ('clients', 'projects') for r in snapshot[key]):
        raise ValueError('Snapshot contains records outside the requested tenant')
    workbook = openpyxl.load_workbook(source, read_only=False, data_only=False)
    cached = openpyxl.load_workbook(source, read_only=False, data_only=True)
    sheet = workbook[f'{year}. évi indexált ár']
    cached_sheet = cached[sheet.title]
    expected_headers = HEADERS.copy()
    expected_headers[4] = f'{year}. évi díjak'
    if [sheet.cell(4, c).value for c in range(1, 12)] != expected_headers:
        raise ValueError('Unexpected worksheet headers')
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    payload, unmatched, skipped = [], [], []
    seen = set()
    for row in range(5, sheet.max_row + 1):
        v = [sheet.cell(row, c).value for c in range(1, 12)]
        review = dict(zip(expected_headers, v))
        review['source_row'] = row
        if not any(value is not None for value in v):
            continue
        if not v[1] or not v[2]:
            skipped.append({**review, 'reason': 'Missing client code or project name'})
            continue
        clients = [c for c in snapshot['clients'] if normalize(c['client_code']) == normalize(v[1])]
        projects = [p for p in snapshot['projects'] if len(clients) == 1
                    and p['client_id'] == clients[0]['id'] and normalize(p['name']) == normalize(v[2])]
        fee = indexed_fee(sheet, cached_sheet, row)
        review[expected_headers[4]] = fee
        if len(clients) != 1 or len(projects) != 1:
            unmatched.append({**review, 'reason': 'No unique client/project name match'})
            continue
        project = projects[0]
        if project['id'] in seen:
            raise ValueError(f'Multiple rows map to project {project["id"]}')
        seen.add(project['id'])
        if not v[5] or not re.fullmatch(r'[A-Z]{3}', str(v[6] or '')):
            raise ValueError(f'Missing unit/currency in row {row}')
        original_text = v[3] if isinstance(v[3], str) else None
        if original_text and original_text.startswith('='):
            raise ValueError(f'Unsupported original-fee formula in D{row}')
        payload.append({
            'project_id': project['id'], 'company_id': company_id, 'pricing_year': year,
            'original_fee': None if original_text else money(v[3]),
            'original_fee_text': original_text, 'indexed_fee': fee,
            'unit': v[5], 'currency': v[6], 'original_unit': v[5], 'original_currency': v[6], 'contract_indexation': v[7],
            'indexation_basis': v[8], 'attachment_note': v[9], 'po_number': v[10],
            'source_file': source.name, 'source_sheet': sheet.title, 'source_row': row,
            'source_sha256': digest, 'source_formula': v[4] if isinstance(v[4], str) else None,
            'client_id': clients[0]['id'], 'expected_project_name': project['name'],
        })
    if not payload:
        raise ValueError('No project prices matched')
    return payload, unmatched, skipped


def seed_sql(payload):
    # Only validated UUIDs, a numeric count, and JSON data enter SQL. JSON stays
    # inside a dollar-quoted literal; reject its delimiter before constructing SQL.
    data = json.dumps(payload, ensure_ascii=False)
    if '$prices$' in data:
        raise ValueError('Source contains the SQL data delimiter')
    columns = [k for k in payload[0] if k not in ('client_id', 'expected_project_name')]
    names = ', '.join(columns)
    selects = ', '.join(f'incoming.{name}' for name in columns)
    company_id = str(UUID(payload[0]['company_id']))
    return f"""begin;
set local lock_timeout = '5s';
do $seed$
declare
  expected jsonb := $prices${data}$prices$::jsonb;
  valid_count integer;
begin
  perform 1 from public.companies
    where id = '{company_id}' and name = 'Ecovis Compliance'
      and is_active and not is_deleted for share;
  if not found then raise exception 'Target Ecovis tenant is unavailable'; end if;
  -- Lock all matched projects; reject a stale snapshot or tenant mismatch.
  perform 1 from public.projects p
    join jsonb_array_elements(expected) e on p.id = (e->>'project_id')::uuid
    for share of p;
  select count(*) into valid_count
    from jsonb_array_elements(expected) e join public.projects p
      on p.id = (e->>'project_id')::uuid
      and p.company_id = '{company_id}'
      and p.company_id = (e->>'company_id')::uuid
      and p.client_id = (e->>'client_id')::uuid
      and p.name = e->>'expected_project_name';
  if valid_count <> {len(payload)} then
    raise exception 'Project mapping changed; regenerate the seed';
  end if;
  insert into public.project_pricing ({names})
    select {selects}
    from jsonb_populate_recordset(null::public.project_pricing, expected) incoming
    on conflict (project_id, pricing_year) do nothing;
  select count(*) into valid_count
    from jsonb_array_elements(expected) e join public.project_pricing stored
      on stored.project_id = (e->>'project_id')::uuid
      and stored.pricing_year = (e->>'pricing_year')::smallint
    where (to_jsonb(stored) - array['created_at','updated_at']) =
      ((e - array['client_id','expected_project_name']) ||
        jsonb_build_object('original_fee', (e->>'original_fee')::numeric,
                           'indexed_fee', (e->>'indexed_fee')::numeric));
  if valid_count <> {len(payload)} then
    raise exception 'Stored prices differ from source; all inserts rolled back';
  end if;
end;
$seed$;
commit;
"""


def write_csv(path, rows):
    if not rows:
        path.write_text('', encoding='utf-8')
        return
    with path.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        # Prevent workbook text from becoming executable CSV formulas.
        writer.writerows({k: "'" + v if isinstance(v, str) and v[:1] in '=+-@' else v
                          for k, v in row.items()} for row in rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--snapshot', type=Path, required=True)
    parser.add_argument('--company-id', required=True)
    parser.add_argument('--year', type=int, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    payload, unmatched, skipped = prepare(args.workbook, json.loads(args.snapshot.read_text()), args.company_id, args.year)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / 'seed.sql').write_text(seed_sql(payload), encoding='utf-8')
    (args.output_dir / 'expected.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    write_csv(args.output_dir / 'matched.csv', payload)
    write_csv(args.output_dir / 'unmatched.csv', unmatched)
    write_csv(args.output_dir / 'skipped.csv', skipped)
    summary = {'company_id': args.company_id, 'pricing_year': args.year,
               'source_sha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
               'matched': len(payload), 'unmatched': len(unmatched), 'skipped': len(skipped),
               'indexed_prices': sum(r['indexed_fee'] is not None for r in payload),
               'unset_indexed_prices': sum(r['indexed_fee'] is None for r in payload),
               'text_original_fees': sum(r['original_fee_text'] is not None for r in payload),
               'currency_counts': dict(Counter(r['currency'] for r in payload))}
    (args.output_dir / 'summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
    print(json.dumps(summary))


if __name__ == '__main__':
    main()
